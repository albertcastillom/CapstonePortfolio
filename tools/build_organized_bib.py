from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("/Users/albertcastillo/Desktop/UWB_Summer26/497/Bib.docx")
OUTPUT = Path("/Users/albertcastillo/repos/CapstonePortfolio/Bib-organized.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5F6B76"
LIGHT_BLUE = "E8EEF5"
LIGHT_GOLD = "FFF4CE"
GOLD = "7A5A00"
RULE = "CBD5E1"


def set_run_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    r_pr.append(sz)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def create_decimal_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "100")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_label_paragraph(doc, label, text, keep_next=False, color=INK):
    p = doc.add_paragraph(style="Source Detail")
    p.paragraph_format.keep_with_next = keep_next
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, size=10.5, color=DARK_BLUE, bold=True)
    text_run = p.add_run(text)
    set_run_font(text_run, size=10.5, color=color)
    return p


def add_callout(doc, label, text, fill, label_color, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "8")
        border.set(qn("w:color"), fill)
        borders.append(border)
    p_pr.append(borders)
    label_run = p.add_run(label)
    set_run_font(label_run, size=10, color=label_color, bold=True)
    text_run = p.add_run(text)
    set_run_font(text_run, size=10, color=INK)
    return p


def add_entry(doc, number, author, title, container, year, links, covers, used, confirm=None):
    p = doc.add_paragraph(style="Source Title")
    p.paragraph_format.keep_with_next = True
    num = p.add_run(f"{number}. ")
    set_run_font(num, size=11, color=BLUE, bold=True)
    author_text = author if author.endswith(".") else author + "."
    author_run = p.add_run(author_text + " ")
    set_run_font(author_run, size=11, bold=True)
    title_run = p.add_run(f'"{title}." ')
    set_run_font(title_run, size=11, italic=True)
    container_run = p.add_run(container)
    set_run_font(container_run, size=11, color=MUTED)
    if year:
        suffix = year if year.endswith(".") else f"{year}."
        year_run = p.add_run(f", {suffix}")
        set_run_font(year_run, size=11, color=MUTED)

    link_p = doc.add_paragraph(style="Source Links")
    link_p.paragraph_format.keep_with_next = True
    for idx, (label, url) in enumerate(links):
        if idx:
            sep = link_p.add_run("  |  ")
            set_run_font(sep, size=10, color=MUTED)
        add_hyperlink(link_p, label, url)

    add_label_paragraph(doc, "What it covers", covers, keep_next=True)
    use_p = add_label_paragraph(doc, "How it supported the project", used, keep_next=bool(confirm))
    use_p.paragraph_format.space_after = Pt(4 if confirm else 10)

    if confirm:
        add_callout(doc, "CONFIRM WITH ALBERT: ", confirm, LIGHT_GOLD, GOLD, after=10)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, base in (
        ("Source Title", "Normal"),
        ("Source Links", "Normal"),
        ("Source Detail", "Normal"),
    ):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.base_style = doc.styles[base]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")

    doc.styles["Source Title"].paragraph_format.space_before = Pt(4)
    doc.styles["Source Title"].paragraph_format.space_after = Pt(2)
    doc.styles["Source Title"].paragraph_format.line_spacing = 1.15
    doc.styles["Source Links"].paragraph_format.space_before = Pt(0)
    doc.styles["Source Links"].paragraph_format.space_after = Pt(4)
    doc.styles["Source Links"].paragraph_format.left_indent = Inches(0.25)
    doc.styles["Source Detail"].paragraph_format.space_before = Pt(0)
    doc.styles["Source Detail"].paragraph_format.space_after = Pt(3)
    doc.styles["Source Detail"].paragraph_format.left_indent = Inches(0.25)
    doc.styles["Source Detail"].paragraph_format.line_spacing = 1.15


def add_title_block(doc):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("CAPSTONE WORKING DRAFT")
    set_run_font(run, size=9, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Organized Annotated Bibliography")
    set_run_font(run, size=24, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("ScreenTimeAppIOS - research, learning, technical decisions, and implementation")
    set_run_font(run, size=12.5, color=MUTED)

    for label, value in (
        ("Prepared for", "Albert Castillo's capstone portfolio"),
        ("Organization", "Grouped by the project's development path"),
        ("Status", "Draft for review before moving entries to the website"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        lr = p.add_run(f"{label}: ")
        set_run_font(lr, size=10.5, color=DARK_BLUE, bold=True)
        vr = p.add_run(value)
        set_run_font(vr, size=10.5, color=INK)

    add_callout(
        doc,
        "Drafting note: ",
        "The annotations below are based on the source topics, the stated learning path "
        "(React Native to Swift/SwiftUI), and a review of the current ScreenTimeAppIOS code. "
        "Yellow notes identify details to confirm before publication.",
        LIGHT_BLUE,
        DARK_BLUE,
        after=12,
    )


def add_questions(doc):
    doc.add_heading("Questions to resolve before website publication", level=1)
    questions = [
        "Realtime versus leaderboard: Was Supabase Realtime first explored for a live leaderboard in an earlier prototype, then reused for focus-request updates? In the current project, Realtime listens for accepted focus requests, while the leaderboard screen still displays placeholder data.",
        "Redis: Did the Redis course influence a planned leaderboard/cache design, or was Redis an option you researched and ultimately did not use?",
        "Broad architecture videos: What specific parts of The Complete App Development Tech Stack and the long SaaS course affected your project - technology selection, database design, authentication, Supabase setup, or something else?",
        "RLS citation: The YouTube channel name currently appears only as 'G'. Do you know the creator's full name or the course/playlist name you want shown?",
        "Chennai study: The saved link is a UW Libraries/EBSCO proxy link. Should the public website keep that institutional link, or should we replace it with a public DOI/PubMed link when available?",
        "Audience emphasis: Should the research section present the app mainly as a tool for students/teens, working adults, or a broader group seeking accountability and focus?",
    ]
    num_id = create_decimal_numbering(doc)
    for question in questions:
        p = doc.add_paragraph()
        apply_numbering(p, num_id)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(question)
        set_run_font(r, size=10.5, color=INK)


def build():
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("ScreenTimeAppIOS | Annotated Bibliography Draft")
    set_run_font(hr, size=9, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    add_page_number(fp)

    add_title_block(doc)

    doc.add_heading("1. Problem research and user needs", level=1)
    p = doc.add_paragraph(
        "These sources establish why screen-time management and focused work matter. "
        "Together they support the capstone's problem statement, user need, and emphasis on accountability rather than restriction alone."
    )
    p.paragraph_format.space_after = Pt(10)

    add_entry(doc, 1, "Monica Anderson, Michelle Faverio, and Eugenie Park",
              "How Teens and Parents Approach Screen Time", "Pew Research Center", "2024",
              [("Open source", "https://www.pewresearch.org/internet/2024/03/11/how-teens-and-parents-approach-screen-time/")],
              "A representative survey of 1,453 U.S. teens ages 13-17 and their parents about phone use, emotions, family conflict, and attempts to reduce screen time.",
              "Provides user-centered evidence that many teens recognize excessive phone use and that families struggle to manage it. This supports the need for a collaborative focus tool and informed the user/problem framing.")

    add_entry(doc, 2, "S. Amuthan et al.",
              "Screen Time and Its Determinants among Adolescents in Chennai: A Cross-Sectional Study",
              "Journal of Family Medicine and Primary Care / EBSCOhost", "2026",
              [("UW Libraries source", "https://research-ebsco-com.offcampus.lib.washington.edu/c/2onyl7/search/details/abvcjpy4zz?db=a9h&limiters=RV%3AY&q=Screen%20time%20and%20focus&searchMode=all")],
              "Examines weekday and weekend screen time among adolescents and relates it to sleep, physical activity, peer interaction, eyesight, and weight-related factors.",
              "Adds health and lifestyle context to the project by showing that screen time is connected to routines beyond the phone itself. It can support discussion of healthier focus habits and the broader stakes of the app.",
              "Confirm that this is the intended article and whether the portfolio should use this UW-proxied link or a public record/DOI.")

    add_entry(doc, 3, "Giorgia Bondanini, Cristina Giovanelli, Nicola Mucci, and Gabriele Giorgi",
              "The Dual Impact of Digital Connectivity: Balancing Productivity and Well-Being in the Modern Workplace",
              "International Journal of Environmental Research and Public Health", "2025",
              [("Open source", "https://www.mdpi.com/1660-4601/22/6/845")],
              "A systematic review of how digital connectivity can improve flexibility and productivity while also contributing to overload, burnout, fatigue, and sleep disruption.",
              "Helps frame the app as a balance tool rather than an anti-technology product. It supports the idea that users benefit from structured offline/focus periods while keeping useful digital connections.")

    add_entry(doc, 4, "Eilish Duke and Christian Montag",
              "Smartphone Addiction, Daily Interruptions and Self-Reported Productivity",
              "Addictive Behaviors Reports", "2017",
              [("Open source", "https://www.sciencedirect.com/science/article/pii/S2352853217300159")],
              "Reports relationships among smartphone overuse, frequent interruptions, hours lost to phone use, and perceived productivity in a sample of 262 participants.",
              "Directly supports the capstone's focus-session concept: reducing interruptions can help users protect concentration and productivity. It provides research grounding for app blocking during a committed task.")

    doc.add_heading("2. Initial React Native approach and architecture exploration", level=1)
    p = doc.add_paragraph(
        "These sources reflect the project's first phase, when React Native and Expo were considered for cross-platform development, and the later feasibility work that led to a native Swift pivot."
    )
    p.paragraph_format.space_after = Pt(10)

    add_entry(doc, 5, "Net Ninja", "Complete React Native Tutorial #1 - Introduction & Setup (Expo)",
              "YouTube", "2025", [("Watch video", "https://www.youtube.com/watch?v=J2j1yk-34OY")],
              "Introduces React Native project setup with Expo and the basic workflow for building a cross-platform mobile interface.",
              "Supported the early React Native learning phase and helped establish the initial project structure before Apple Screen Time requirements pushed the implementation toward native Swift.")

    add_entry(doc, 6, "Expo", "How to Create a Native Module with the Expo Modules API",
              "YouTube", "2024", [("Watch video (saved at 14:45)", "https://www.youtube.com/watch?v=CdaQSlyGik8&t=885s")],
              "Shows how Expo/React Native code can call platform-specific native functionality through a custom native module.",
              "Was relevant when evaluating whether Apple's native-only Screen Time frameworks could be bridged into the original React Native app. The added complexity and performance/maintenance tradeoffs helped justify the full Swift/SwiftUI pivot.")

    add_entry(doc, 7, "Warren Day", "The Complete App Development Tech Stack", "YouTube", "2026",
              [("Watch video", "https://www.youtube.com/watch?v=zinyWfgNQgU")],
              "Surveys the major layers and technology choices involved in building a modern application.",
              "Likely supported early architecture and stack selection, including the comparison between cross-platform and native development.",
              "Which specific decision did this source influence: React Native versus Swift, backend selection, deployment, or the overall architecture?")

    add_entry(doc, 8, "JavaScript Mastery",
              "SaaS App Full Course 2026: Launch Your SaaS in Under 7 Days with Next.js, Supabase & Payments",
              "YouTube", "2025",
              [("Watch video (saved at 3:48:39)", "https://www.youtube.com/watch?v=XUkNR-JfHwo&t=13719s")],
              "Builds a full application using a JavaScript frontend and Supabase-backed services.",
              "Appears to have supported the project's Supabase/backend learning during the React-oriented phase, but the exact saved section could not be verified from available metadata.",
              "What was happening around 3:48:39, and did you use it for Supabase setup, authentication, database work, or another feature?")

    doc.add_heading("3. Realtime concepts, Supabase, and backend security", level=1)
    p = doc.add_paragraph(
        "These resources support the app's shared backend: authentication, profiles and friendships, focus-request data, secure database access, Edge Functions, and live updates."
    )
    p.paragraph_format.space_after = Pt(10)

    add_entry(doc, 9, "freeCodeCamp.org", "Redis Course - In-Memory Database Tutorial", "YouTube", "2020",
              [("Watch video", "https://www.youtube.com/watch?v=XCsS_NVAa1g")],
              "Explains Redis as an in-memory data store used for fast reads, caching, counters, and realtime-oriented workloads.",
              "May have informed early leaderboard or performance research, but Redis does not appear in the current ScreenTimeAppIOS dependencies or code.",
              "Was Redis considered for leaderboard rankings/caching, used in an earlier prototype, or simply researched and not adopted?")

    add_entry(doc, 10, "freeCodeCamp.org", "A Beginner's Guide to WebSockets", "YouTube", "2018",
              [("Watch video", "https://www.youtube.com/watch?v=8ARodQ4Wlf4")],
              "Introduces persistent, two-way WebSocket communication and how it differs from request-response HTTP.",
              "Built the conceptual foundation for live updates. The final app uses Supabase Realtime channels for accepted focus-request events and includes reconnection plus database reconciliation logic.",
              "You mentioned a realtime leaderboard, but the current Realtime listener handles focus requests and the leaderboard is still static. Was the live leaderboard built in an earlier version or was it the original use case for this research?")

    add_entry(doc, 11, "Supabase", "Getting Started with Supabase Database", "YouTube", "2026",
              [("Watch video", "https://www.youtube.com/watch?v=C9kJwxhdw9A")],
              "Covers Supabase database concepts and the basic workflow for storing and retrieving application data.",
              "Supported the PostgreSQL-backed data layer used for profiles, friendships, and focus requests in the final app.")

    add_entry(doc, 12, "Supabase", "Getting Started with Supabase Edge Functions", "YouTube", "2026",
              [("Watch video", "https://www.youtube.com/watch?v=EQ5h3NSWMbk&list=PL5S4mPUpp4OsWK_UHmQK41DEgqefYeTPN&index=7")],
              "Introduces server-side Edge Functions in the Supabase platform.",
              "Supported the secure account-deletion flow. The final project includes a delete-account Edge Function so privileged deletion logic runs on the server rather than inside the iOS client.")

    add_entry(doc, 13, "Supabase", "Invokes a Supabase Edge Function (Swift Reference)",
              "Supabase Docs", "n.d.", [("Open documentation", "https://supabase.com/docs/reference/swift/functions-invoke")],
              "Documents the Swift API for invoking an Edge Function, decoding responses, passing bodies/headers, and handling function errors.",
              "Maps directly to SupabaseAuthService, where the iOS app calls the delete-account function and decodes the response.")

    add_entry(doc, 14, "G (channel name as displayed)", "Row Level Security with Supabase - Course Part 10 (2025)",
              "YouTube", "2025", [("Watch video", "https://www.youtube.com/watch?v=Vx1q8Nfp0BE")],
              "Explains Supabase/PostgreSQL Row Level Security policies for restricting which rows authenticated users can read or modify.",
              "Informed the security model for user-owned and relationship-based data such as profiles, friendships, and focus requests.",
              "The channel's public display name appears only as 'G'. If you know the full creator or course name, provide it for a cleaner citation.")

    add_entry(doc, 15, "AppStuff", "Supabase Auth with SwiftUI - Done Right (Production Setup Guide)",
              "YouTube", "2026", [("Watch video", "https://www.youtube.com/watch?v=K3OknZU0Lcs")],
              "Demonstrates a production-oriented approach to connecting Supabase authentication with SwiftUI state and navigation.",
              "Supported the final sign-up, login, logout, session, and authentication-state architecture implemented through SupabaseAuthService and AuthManager.")

    add_entry(doc, 16, "Supabase", "Use Supabase with iOS and SwiftUI", "Supabase Docs", "n.d.",
              [("Open quickstart", "https://supabase.com/docs/guides/getting-started/quickstarts/ios-swiftui")],
              "Official quickstart for installing the Swift client, creating a shared client, and reading/writing Supabase data from a SwiftUI app.",
              "Provided the baseline for integrating the Supabase Swift package and building the shared client/service layer used throughout the final app.")

    doc.add_heading("4. Learning Swift and SwiftUI after the pivot", level=1)
    p = doc.add_paragraph(
        "After determining that Apple's Screen Time frameworks were best implemented natively, these resources supported the transition from React-based development to Swift and SwiftUI."
    )
    p.paragraph_format.space_after = Pt(10)

    add_entry(doc, 17, "Swift.org", "Build an iOS App with SwiftUI", "Swift.org", "n.d.",
              [("Open tutorial", "https://www.swift.org/getting-started/swiftui/")],
              "Official beginner tutorial covering Xcode project creation, the SwiftUI view model, stacks, text, images, buttons, modifiers, and program state.",
              "Helped establish the basic SwiftUI mental model needed to rebuild the app natively and create the initial views and interactions.")

    add_entry(doc, 18, "Sean Allen", "SwiftUI Fundamentals - Full Course", "YouTube", "2023",
              [("Watch video", "https://www.youtube.com/watch?v=b1oC7sLIgpI")],
              "A broad beginner course on SwiftUI layout, data flow, controls, navigation, and common application patterns.",
              "Served as a structured foundation for learning SwiftUI while developing the final app's home, friends, profile, authentication, and leaderboard screens.")

    add_entry(doc, 19, "Paul Hudson", "WeSplit SwiftUI Tutorial Series: App Structure, Forms, Program State, and Bindings",
              "YouTube", "2023",
              [("App structure (1/11)", "https://www.youtube.com/watch?v=sZSlTDlo0Ag"),
               ("Creating a form (2/11)", "https://www.youtube.com/watch?v=4Ui09XbYf1A"),
               ("Program state (4/11)", "https://www.youtube.com/watch?v=X0Rw3uc3kj8"),
               ("Bindings (5/11)", "https://www.youtube.com/watch?v=wMIHdsFYGC0")],
              "A focused sequence explaining the structure of a SwiftUI app, form construction, state mutation, and two-way bindings between data and controls.",
              "Supported the transition from React state concepts to SwiftUI's declarative data flow. These lessons apply directly to forms, pickers, filters, login/sign-up fields, and screen updates across the app.")

    add_entry(doc, 20, "freeCodeCamp.org", "SwiftUI Course for Beginners - Create an iOS App from Scratch",
              "YouTube", "2025", [("Watch video (saved at 20:19)", "https://www.youtube.com/watch?v=-VC3hIEL7eQ&t=1219s")],
              "A project-based introduction to creating an iOS application with SwiftUI.",
              "Reinforced the SwiftUI project workflow and helped translate beginner concepts into a complete application structure during the native rebuild.")

    doc.add_heading("5. Apple Screen Time frameworks and app blocking", level=1)
    p = doc.add_paragraph(
        "These were the central implementation resources for the capstone's defining feature: requesting Screen Time authorization, selecting apps, scheduling focus sessions, and applying/removing shields while preserving Apple's privacy model."
    )
    p.paragraph_format.space_after = Pt(10)

    add_entry(doc, 21, "Apple", "Screen Time Technology Frameworks", "Apple Developer Documentation", "n.d.",
              [("Open documentation", "https://developer.apple.com/documentation/screentimeapidocumentation")],
              "Official entry point for Apple's Screen Time technologies, including FamilyControls, ManagedSettings, and DeviceActivity.",
              "Was the authoritative reference for capabilities, privacy constraints, permissions, framework responsibilities, and the native architecture used by the final app.")

    add_entry(doc, 22, "Julius Brussee",
              "A Developer's Guide to Apple's Screen Time APIs (FamilyControls, ManagedSettings, DeviceActivity)",
              "Medium", "2025",
              [("Open article", "https://medium.com/@juliusbrussee/a-developers-guide-to-apple-s-screen-time-apis-familycontrols-managedsettings-deviceactivity-e660147367d7")],
              "Explains how the three main frameworks work together: FamilyControls for authorization and selection, ManagedSettings for enforcement, and DeviceActivity for scheduling and monitoring.",
              "Helped turn Apple's separate framework documentation into a practical implementation plan. The final project follows this pattern through PermissionManager, AppBlockingModel, RestrictionsService, and the DeviceActivity monitor extension.")

    add_entry(doc, 23, "Ezgi Ustunel", "Screen Time API", "iOS Nest / Medium", "2024",
              [("Open article", "https://medium.com/ios-nest/screen-time-api-d1110751d2ce")],
              "Provides code examples for FamilyControls authorization, the FamilyActivityPicker, saving selections, DeviceActivity schedules/events, and monitor callbacks.",
              "Supported hands-on implementation and debugging of permission requests, app selection persistence, monitoring schedules, and the extension that activates or clears restrictions.")

    add_questions(doc)

    doc.core_properties.title = "Organized Annotated Bibliography - ScreenTimeAppIOS"
    doc.core_properties.subject = "Capstone research and development sources"
    doc.core_properties.keywords = "ScreenTimeAppIOS, capstone, annotated bibliography, SwiftUI, Supabase, Screen Time"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
